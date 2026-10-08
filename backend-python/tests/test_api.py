import pytest
import io
import fitz # PyMuPDF
from PIL import Image
from docx import Document
from fastapi.testclient import TestClient
from app.main import app

client = TestClient(app)

def create_dummy_image(format_name="JPEG") -> bytes:
    img = Image.new("RGB", (100, 100), color="blue")
    buf = io.BytesIO()
    img.save(buf, format=format_name)
    return buf.getvalue()

def create_dummy_pdf() -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), "AllTools Platform Unit Test PDF Document 2026")
    pdf_bytes = doc.write()
    doc.close()
    return pdf_bytes

def create_dummy_docx() -> bytes:
    doc = Document()
    doc.add_heading("AllTools Document Conversion Test", level=1)
    doc.add_paragraph("This is a paragraph converted from DOCX to PDF format.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

# 1. Health & Status
def test_health_check():
    res = client.get("/api/health")
    assert res.status_code == 200
    assert res.json()["status"] == "healthy"

# 2. Password Generator
def test_password_generator():
    res = client.get("/api/generator/password?length=20&uppercase=true&numbers=true&symbols=true")
    assert res.status_code == 200
    assert len(res.json()["password"]) == 20

# 3. UUID Generator
def test_uuid_generator():
    res = client.get("/api/generator/uuid?count=3")
    assert res.status_code == 200
    assert len(res.json()["uuids"]) == 3

# 4. QR Generator
def test_qr_generator():
    res = client.post("/api/generator/qr", json={"data": "https://example.com", "format": "png"})
    assert res.status_code == 200
    assert res.json()["success"] is True

# 5. Fake Address & Mock Data
def test_fake_address_data():
    res = client.get("/api/test-data/address?count=3")
    assert res.status_code == 200
    assert len(res.json()["data"]) == 3

# 6. Hash Generator
def test_hash_generator():
    res = client.post("/api/dev/hash", json={"text": "AllTools 2026"})
    assert res.status_code == 200
    data = res.json()
    assert data["success"] is True
    assert "sha256" in data["hashes"]

# 7. Image Converter (Multi-way)
def test_image_converter():
    img_bytes = create_dummy_image("JPEG")
    files = {"file": ("test.jpg", img_bytes, "image/jpeg")}
    res = client.post("/api/image/convert", files=files, data={"target_format": "png"})
    assert res.status_code == 200
    assert len(res.content) > 0

# 8. Remove Background (Standard)
def test_remove_background():
    img_bytes = create_dummy_image("JPEG")
    files = {"file": ("test.jpg", img_bytes, "image/jpeg")}
    res = client.post("/api/image/remove-background", files=files)
    assert res.status_code == 200
    assert len(res.content) > 0

# 9. Remove Logo BG (Smooth)
def test_remove_logo_background():
    img_bytes = create_dummy_image("PNG")
    files = {"file": ("logo.png", img_bytes, "image/png")}
    res = client.post("/api/image/remove-logo-bg", files=files, data={"smoothness": "3", "tolerance": "25"})
    assert res.status_code == 200
    assert len(res.content) > 0

# 10. Image to ICO (Favicon)
def test_image_to_ico():
    img_bytes = create_dummy_image("PNG")
    files = {"file": ("logo.png", img_bytes, "image/png")}
    res = client.post("/api/image/to-ico", files=files)
    assert res.status_code == 200
    assert len(res.content) > 0

# 11. Image to SVG Vector
def test_image_to_svg():
    img_bytes = create_dummy_image("PNG")
    files = {"file": ("logo.png", img_bytes, "image/png")}
    res = client.post("/api/image/to-svg", files=files)
    assert res.status_code == 200
    assert "<svg" in res.text

# 12. PDF to Word (DOCX)
def test_pdf_to_word():
    pdf_bytes = create_dummy_pdf()
    files = {"file": ("document.pdf", pdf_bytes, "application/pdf")}
    res = client.post("/api/pdf/pdf-to-word", files=files)
    assert res.status_code == 200
    assert len(res.content) > 0

# 13. Word (DOCX) to PDF
def test_word_to_pdf():
    docx_bytes = create_dummy_docx()
    files = {"file": ("sample.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    res = client.post("/api/doc/word-to-pdf", files=files)
    assert res.status_code == 200
    assert len(res.content) > 0

# 14. PDF to Text Extractor
def test_pdf_to_text():
    pdf_bytes = create_dummy_pdf()
    files = {"file": ("doc.pdf", pdf_bytes, "application/pdf")}
    res = client.post("/api/pdf/pdf-to-text", files=files)
    assert res.status_code == 200
    assert "AllTools Platform Unit Test" in res.json()["text"]

# 15. Merge PDFs
def test_merge_pdfs():
    pdf1 = create_dummy_pdf()
    pdf2 = create_dummy_pdf()
    files = [
        ("files", ("doc1.pdf", pdf1, "application/pdf")),
        ("files", ("doc2.pdf", pdf2, "application/pdf"))
    ]
    res = client.post("/api/pdf/merge", files=files)
    assert res.status_code == 200
    assert len(res.content) > 0
